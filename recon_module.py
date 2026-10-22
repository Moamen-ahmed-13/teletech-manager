# -*- coding: utf-8 -*-
"""
recon_module.py - BOQ vs. MIR reconciliation for TELETECH CAD UI.

Design rules (see README_RECON.md):
  * No global state: every function receives the project directory explicitly,
    so data of different projects/users can never mix.
  * OCR output is only a DRAFT. It is stored untouched in rec['ocr']; the human
    reviewer's values live in rec['user']. Only records that are reviewed AND
    have an approved BOQ match AND an accepted MIR status are counted.
  * Nothing is silently capped, defaulted to zero, or hidden: overages,
    negative remaining quantities, missing prices and unit conflicts are
    reported as flags/exceptions and as "unavailable" values.
"""
import difflib, hashlib, io, json, os, re, shutil, subprocess, tempfile, threading, time, uuid
from contextlib import contextmanager
from pathlib import Path

MAX_UPLOAD_BYTES = 100 * 1024 * 1024
MAX_PDF_PAGES = 80
COUNT_STATUSES_DEFAULT = ('A', 'B')
MIR_STATUS_LABELS = {'A': 'Approved', 'B': 'Approved with note', 'C': 'Not approved', 'D': 'Rejected'}
DEFINITION_NOTE = ('A MIR (Material Inspection Request) records material submitted for inspection/approval. '
                   'Confirmed MIR quantity is NOT proof of delivery, installation, consumption, or invoicing. '
                   'Quantities are counted only for reviewed records with an approved BOQ match and MIR status A or B; '
                   'statuses C/D and unknown are reported separately.')


class ReconError(ValueError):
    """User-facing validation error (HTTP 400)."""


# ----------------------------------------------------------------------------
# storage (per project, atomic, locked)
# ----------------------------------------------------------------------------
_LOCKS, _LOCKS_GUARD = {}, threading.Lock()
_PROC_ID = uuid.uuid4().hex  # identifies this server process (see load_state / start_ocr)


def _plock(project_dir):
    k = str(Path(project_dir).resolve())
    with _LOCKS_GUARD:
        return _LOCKS.setdefault(k, threading.RLock())


def _now():
    return time.strftime('%Y-%m-%dT%H:%M:%SZ', time.gmtime())


def recon_dir(project_dir):
    d = Path(project_dir) / 'reconciliation'
    for sub in ('', 'boq', 'mir', 'pages', 'exports'):
        (d / sub).mkdir(parents=True, exist_ok=True)
    return d


def _new_state():
    return {'version': 1, 'created_at': _now(), 'updated_at': None,
            'settings': {'currency_label': '', 'count_statuses': list(COUNT_STATUSES_DEFAULT)},
            'boq': None, 'boq_history': [], 'boq_overrides': {}, 'mir_docs': [], 'records': [], 'audit': []}


def load_state(project_dir):
    p = recon_dir(project_dir) / 'state.json'
    if not p.exists():
        return _new_state()
    try:
        data = json.loads(p.read_text(encoding='utf-8'))
    except (OSError, ValueError) as e:
        # never silently discard a corrupt file: keep it for inspection
        bad = p.with_name('state.corrupt.%s.json' % int(time.time()))
        try: p.replace(bad)
        except OSError: pass
        raise ReconError('Reconciliation data file was unreadable and has been moved to %s (%s).' % (bad.name, e))
    base = _new_state()
    for k, v in base.items():
        data.setdefault(k, v)
    for k, v in base['settings'].items():
        data['settings'].setdefault(k, v)
    for d in data.get('mir_docs', []):
        # an OCR job that belonged to a previous server process can never finish: say so instead of showing "running" forever
        if d.get('ocr_status') == 'running' and d.get('ocr_proc') != _PROC_ID:
            d['ocr_status'] = 'interrupted'
    return data


def save_state(project_dir, st):
    d = recon_dir(project_dir)
    st['updated_at'] = _now()
    tmp = d / ('state.%s.tmp' % uuid.uuid4().hex)
    tmp.write_text(json.dumps(st, ensure_ascii=False, indent=1), encoding='utf-8')
    tmp.replace(d / 'state.json')


@contextmanager
def locked_state(project_dir, save=True):
    with _plock(project_dir):
        st = load_state(project_dir)
        yield st
        if save:
            save_state(project_dir, st)


def _audit(st, user, action, detail=''):
    st['audit'].append({'ts': _now(), 'user': (user or {}).get('username', ''), 'action': action, 'detail': str(detail)[:500]})
    if len(st['audit']) > 5000:
        st['audit'] = st['audit'][-5000:]


# ----------------------------------------------------------------------------
# helpers: numbers, text, units, brands
# ----------------------------------------------------------------------------
_AR_DIGITS = str.maketrans('٠١٢٣٤٥٦٧٨٩۰۱۲۳۴۵۶۷۸۹', '01234567890123456789')


def to_number(v):
    """Return float or None. Never guesses: non-numeric text -> None."""
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v) if v == v else None
    s = str(v).translate(_AR_DIGITS).strip().replace(',', '').replace('\u066b', '.')
    if not s or not re.fullmatch(r'[-+]?\d+(\.\d+)?', s):
        return None
    return float(s)


def _txt(v):
    if v is None:
        return ''
    if isinstance(v, float) and v == int(v):
        v = int(v)
    return re.sub(r'\s+', ' ', str(v).replace('\n', ' ')).strip()


def norm_alnum(s):
    return re.sub(r'[^a-z0-9\u0600-\u06ff]', '', str(s or '').lower())


UNIT_ALIASES = {'pc': 'PC', 'pcs': 'PC', 'piece': 'PC', 'pieces': 'PC', 'ea': 'PC', 'each': 'PC', 'no': 'PC', 'nos': 'PC',
                'roll': 'ROLL', 'rolls': 'ROLL', 'm': 'M', 'mtr': 'M', 'meter': 'M', 'metre': 'M', 'meters': 'M',
                'set': 'SET', 'sets': 'SET', 'lot': 'LOT', 'box': 'BOX', 'boxes': 'BOX', 'ls': 'LS', 'unit': 'UNIT', 'units': 'UNIT'}


def canon_unit(u):
    s = re.sub(r'[^a-z\u0600-\u06ff]', '', str(u or '').lower())
    if not s:
        return None
    return UNIT_ALIASES.get(s, s.upper())


def fuzzy_unit(u):
    """OCR-only helper: 'Rol' -> ROLL. Returns (unit, was_fuzzy)."""
    s = re.sub(r'[^a-z]', '', str(u or '').lower())
    if s in UNIT_ALIASES:
        return UNIT_ALIASES[s], False
    if len(s) >= 3:
        m = difflib.get_close_matches(s, list(UNIT_ALIASES), n=1, cutoff=0.74)
        if m:
            return UNIT_ALIASES[m[0]], True
    return None, False


BRAND_ALIASES = {'comscpe': 'commscope', 'commscpe': 'commscope', 'comscope': 'commscope', 'commscope': 'commscope'}


def brand_key(b):
    s = norm_alnum(b)
    return BRAND_ALIASES.get(s, s)


STOP = {'the', 'and', 'with', 'for', 'of', 'accessories', 'accessory', 'sample', 'part', 'pc', 'pcs', 'type', 'new', 'item'}


def tokens(s):
    s = str(s or '').lower()
    s = re.sub(r'cat\s*[-]?\s*(\d)', r'cat\1', s)
    out = []
    for t in re.findall(r'[a-z0-9\u0600-\u06ff]+', s):
        if t in STOP or len(t) < 2:
            continue
        if t.isdigit() and len(t) < 3:
            continue
        if len(t) > 4 and t.endswith('s') and not t.endswith('ss'):
            t = t[:-1]
        out.append(t)
    return out


def _tok_eq(a, b):
    if a == b:
        return True
    if len(a) >= 5 and len(b) >= 5:
        return difflib.SequenceMatcher(None, a, b).ratio() >= 0.85
    return False


def _hits(src, pool):
    return sum(1 for t in src if any(_tok_eq(t, p) for p in pool))


GENERIC_MODELS = {'cat5e', 'cat6', 'cat6a', 'cat7', 'cat7a', 'om3', 'om4', 'os2', 'rj45', 'rj11', 'ip65', 'ip66', 'poe'}


def model_tokens(s):
    """Model-like tokens: contain a digit and a letter, length >= 4 (e.g. MGS600, M10LE-262 -> m10le)."""
    out = set()
    for t in re.findall(r'[A-Za-z0-9]+', str(s or '')):
        tl = t.lower()
        if len(tl) >= 4 and re.search(r'\d', tl) and re.search(r'[a-z]', tl) and tl not in GENERIC_MODELS:
            out.add(tl)
    return out


def _similar_desc(a, b):
    """True when two MIR descriptions name the same material (tolerates OCR typos such as commscpe/commscope)."""
    ta, tb = tokens(a), tokens(b)
    if not ta or not tb:
        return False
    return min(_hits(ta, tb) / len(ta), _hits(tb, ta) / len(tb)) >= 0.75


# ----------------------------------------------------------------------------
# file validation / upload
# ----------------------------------------------------------------------------
def _sha256(b):
    return hashlib.sha256(b).hexdigest()


def pdf_page_count(path):
    try:
        import fitz
        with fitz.open(str(path)) as d:
            return d.page_count
    except ImportError:
        pass
    except Exception as e:
        raise ReconError('The PDF could not be opened: %s' % e)
    try:
        from pypdf import PdfReader
        return len(PdfReader(str(path)).pages)
    except Exception:
        pass
    try:
        out = subprocess.run(['pdfinfo', str(path)], capture_output=True, text=True, timeout=30).stdout
        m = re.search(r'Pages:\s+(\d+)', out)
        if m:
            return int(m.group(1))
    except Exception:
        pass
    raise ReconError('The PDF could not be opened (no PDF library available or the file is corrupt).')


def save_upload(project_dir, st, kind, filename, payload, user, safe_name_fn, extra=None):
    """Validate + store an original file unchanged. kind: 'boq' | 'mir'."""
    if kind not in ('boq', 'mir'):
        raise ReconError('Unknown upload kind.')
    if not payload:
        raise ReconError('The file is empty.')
    if len(payload) > MAX_UPLOAD_BYTES:
        raise ReconError('File exceeds the 100 MB limit.')
    safe = safe_name_fn(filename)
    ext = Path(safe).suffix.lower()
    if kind == 'boq':
        if ext not in ('.xlsx', '.xlsm'):
            raise ReconError('BOQ must be an .xlsx or .xlsm workbook (save older .xls files as .xlsx first).')
        if not payload.startswith(b'PK\x03\x04'):
            raise ReconError('The file content is not a valid Excel workbook.')
    else:
        if ext != '.pdf':
            raise ReconError('MIR files must be PDF.')
        if not payload[:1024].lstrip().startswith(b'%PDF-'):
            raise ReconError('The file content is not a valid PDF.')
    sha = _sha256(payload)
    known = [d for d in st['mir_docs']] + ([st['boq']['file']] if st.get('boq') else []) + list(st.get('boq_history', []))
    for d in known:
        if d.get('sha256') == sha:
            raise ReconError('This exact file was already uploaded to this project as "%s". Duplicate not added.' % d.get('filename'))
    rd = recon_dir(project_dir)
    fid = uuid.uuid4().hex
    stored = fid + ext
    target = rd / ('boq' if kind == 'boq' else 'mir') / stored
    with target.open('xb') as f:
        f.write(payload)
    meta = {'id': fid, 'filename': safe, 'stored_name': stored, 'sha256': sha, 'size_bytes': len(payload),
            'uploaded_by': (user or {}).get('username', ''), 'uploaded_at': _now()}
    try:
        if kind == 'mir':
            n = pdf_page_count(target)
            if n < 1 or n > MAX_PDF_PAGES:
                raise ReconError('PDF must have between 1 and %d pages (this one has %d).' % (MAX_PDF_PAGES, n))
            meta.update({'pages': n, 'ocr_status': 'not_run', 'ocr_backend': None, 'ocr_at': None,
                         'doc_ref_no_hint': _ref_from_filename(safe)})
            st['mir_docs'].append(meta)
            _audit(st, user, 'MIR_UPLOAD', '%s pages=%d sha=%s' % (safe, n, sha[:12]))
        else:
            parsed = parse_boq(target, display_name=safe)
            if st.get('boq'):
                st['boq_history'].append(st['boq']['file'])
            meta['source_note'] = (extra or {}).get('source_note', '')
            parsed['file'] = meta
            st['boq'] = parsed
            _revalidate_matches(st)
            _audit(st, user, 'BOQ_UPLOAD', '%s items=%d sheet=%s' % (safe, len(parsed['items']), parsed['sheet']))
    except Exception:
        try: target.unlink()
        except OSError: pass
        raise
    return meta


def _ref_from_filename(name):
    m = re.search(r'MIR[\s_#\-]*0*(\d{1,4})', name, re.I)
    return 'MIR-%03d' % int(m.group(1)) if m else ''


def stored_path(project_dir, kind, meta):
    base = (recon_dir(project_dir) / ('boq' if kind == 'boq' else 'mir')).resolve()
    p = (base / str(meta['stored_name'])).resolve()
    if p.parent != base or not p.is_file():
        raise ReconError('Stored file is missing.')
    return p


# ----------------------------------------------------------------------------
# BOQ parsing
# ----------------------------------------------------------------------------
HEADER_SYNS = {
    'sn': {'s.n', 'sn', 's n', 'no', 'no.', '#', 'item no', 'item no.', 'م'},
    'item': {'item', 'description', 'item description', 'البند', 'بيان الاعمال'},
    'reference': {'reference', 'ref', 'ref.', 'spec', 'specification', 'specifications'},
    'part': {'part #', 'part#', 'part no', 'part no.', 'part number', 'part', 'model', 'p/n', 'pn'},
    'qty': {'qty', 'qty.', 'quantity', 'الكمية'},
    'unit': {'unit', 'uom', 'unit of measure', 'الوحدة'},
    'brand': {'brand', 'manufacturer', 'make'},
    'unit_price': {'unit price', 'unit rate', 'rate', 'price', 'سعر الوحدة'},
    'total_price': {'total price', 'total', 'amount', 'total value', 'الاجمالي', 'الاجمالى'},
}


def _hnorm(v):
    return re.sub(r'\s+', ' ', str(v or '').strip().lower())


def _map_header(values):
    cmap, dups = {}, []
    for i, v in enumerate(values):
        h = _hnorm(v)
        if not h:
            continue
        for field, syns in HEADER_SYNS.items():
            if h in syns:
                if field in cmap:
                    dups.append('%s (column %d ignored)' % (field, i + 1))
                else:
                    cmap[field] = i
                break
    return cmap, dups


def _sheet_candidate(ws):
    best = None
    for r_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=min(ws.max_row, 40), values_only=True), 1):
        cmap, dups = _map_header(row)
        if 'qty' in cmap and ('item' in cmap or 'reference' in cmap or 'part' in cmap) and len(cmap) >= 3:
            score = len(cmap) + (3 if 'unit_price' in cmap else 0) + (1 if 'total_price' in cmap else 0)
            if best is None or score > best['score']:
                best = {'row': r_idx, 'map': cmap, 'dups': dups, 'score': score}
    return best


def col_letter(i):
    s, i = '', i + 1
    while i:
        i, r = divmod(i - 1, 26)
        s = chr(65 + r) + s
    return s


def parse_boq(path, display_name=None, sheet=None):
    import openpyxl
    try:
        wb = openpyxl.load_workbook(str(path), data_only=True, read_only=False)
    except Exception as e:
        raise ReconError('The workbook could not be read: %s' % e)
    considered, chosen = [], None
    for ws in wb.worksheets:
        c = _sheet_candidate(ws)
        considered.append({'sheet': ws.title, 'dimensions': ws.dimensions,
                           'header_row': c['row'] if c else None, 'score': c['score'] if c else 0,
                           'fields': sorted(c['map']) if c else []})
        if c and (sheet == ws.title or (sheet is None and (chosen is None or c['score'] > chosen[1]['score']))):
            chosen = (ws, c)
    if not chosen:
        raise ReconError('No worksheet with a recognizable BOQ header row (needs Qty plus Item/Reference/Part) was found.')
    ws, c = chosen
    cmap = c['map']
    warnings = []
    for d in c['dups']:
        warnings.append({'severity': 'info', 'type': 'duplicate_header', 'message': 'Header repeated: %s' % d})
    for f in ('unit_price', 'total_price', 'brand', 'unit', 'part'):
        if f not in cmap:
            warnings.append({'severity': 'warning' if f in ('unit_price', 'brand') else 'info', 'type': 'column_missing',
                             'message': 'No "%s" column recognized in sheet "%s".' % (f, ws.title)})
    # cached-value check (formulas without computed values would read as empty)
    try:
        wbf = openpyxl.load_workbook(str(path), data_only=False)
        wsf = wbf[ws.title]
        nofv = 0
        for row in wsf.iter_rows(min_row=c['row'] + 1, max_row=wsf.max_row):
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith('='):
                    if ws.cell(row=cell.row, column=cell.column).value is None:
                        nofv += 1
        if nofv:
            warnings.append({'severity': 'warning', 'type': 'formula_without_value',
                             'message': '%d formula cells have no cached value; open and save the workbook in Excel, then re-upload.' % nofv})
    except Exception:
        pass
    items, seen_part = [], {}
    for r_idx, row in enumerate(ws.iter_rows(min_row=c['row'] + 1, max_row=ws.max_row, values_only=True), c['row'] + 1):
        def g(f):
            i = cmap.get(f)
            return row[i] if i is not None and i < len(row) else None
        item, ref, part, brand = _txt(g('item')), _txt(g('reference')), _txt(g('part')), _txt(g('brand'))
        if not (item or ref or part or brand):
            continue  # blank row or totals row (numbers only): never an item
        qty_raw = g('qty')
        up_raw, tp_raw = g('unit_price'), g('total_price')
        qty, up, tp = to_number(qty_raw), to_number(up_raw), to_number(tp_raw)
        rid = '%s!R%d' % (ws.title, r_idx)
        flags = []
        if g('sn') in (None, ''):
            flags.append('sn_missing')
        if not (item or ref):
            flags.append('description_missing')
        if not brand:
            flags.append('brand_missing')
        if qty is None:
            flags.append('qty_invalid')
        if 'unit_price' in cmap and up is None:
            flags.append('price_missing')
        elif 'unit_price' not in cmap:
            flags.append('price_missing')
        elif up is not None and up < 0:
            flags.append('price_negative')
        if qty is not None and qty < 0:
            flags.append('qty_negative')
        if qty is not None and up is not None and tp is not None and abs(qty * up - tp) > max(0.5, abs(tp) * 0.001):
            flags.append('total_mismatch')
        if part:
            seen_part.setdefault(part, []).append(rid)
        unit_cell = _txt(g('unit'))
        items.append({'id': rid, 'sheet': ws.title, 'row': r_idx, 'sn': _txt(g('sn')), 'item': item, 'reference': ref,
                      'part': part, 'brand': brand, 'qty': qty, 'qty_raw': _txt(qty_raw) if qty is None and qty_raw not in (None, '') else '',
                      'unit': unit_cell or None, 'unit_price': up, 'total_price': tp, 'flags': flags})
    for part, ids in seen_part.items():
        if len(ids) > 1:
            for it in items:
                if it['id'] in ids:
                    it['flags'].append('part_repeated')
            warnings.append({'severity': 'info', 'type': 'part_repeated',
                             'message': 'Part # %s appears on %d BOQ rows (%s); MIR lines for it will be ambiguous until you choose a row.' % (part, len(ids), ', '.join(ids))})
    for it in items:
        for fl in it['flags']:
            if fl in ('qty_invalid', 'total_mismatch', 'price_negative', 'qty_negative', 'description_missing'):
                warnings.append({'severity': 'warning', 'type': fl, 'ref': it['id'], 'message': '%s: %s' % (it['id'], fl.replace('_', ' '))})
    if not items:
        raise ReconError('The BOQ sheet "%s" has a header but no item rows.' % ws.title)
    columns = {f: '%s (%s)' % (col_letter(i), _txt(ws.cell(row=c['row'], column=i + 1).value)) for f, i in cmap.items()}
    return {'filename': display_name, 'sheet': ws.title, 'header_row': c['row'], 'columns': columns, 'items': items,
            'warnings': warnings, 'sheets_considered': considered, 'parsed_at': _now()}


def boq_index(st):
    return {it['id']: it for it in (st.get('boq') or {}).get('items', [])}


def boq_unit(st, it):
    ov = st.get('boq_overrides', {}).get(it['id'], {})
    if ov.get('unit'):
        return canon_unit(ov['unit']), 'user'
    if it.get('unit'):
        return canon_unit(it['unit']), 'sheet'
    return None, None


def _revalidate_matches(st):
    idx = boq_index(st)
    for r in st['records']:
        m = r.get('match') or {}
        if m.get('boq_id') and m['boq_id'] not in idx:
            r['match'] = {'state': 'needs_rematch', 'boq_id': None, 'candidates': [], 'method': 'auto',
                          'note': 'Previously matched BOQ row %s no longer exists in the active BOQ.' % m['boq_id']}


# ----------------------------------------------------------------------------
# page rendering + OCR (draft only)
# ----------------------------------------------------------------------------
def tesseract_cmd():
    c = os.environ.get('TESSERACT_CMD', '').strip()
    if c and Path(c).is_file():
        return c
    w = shutil.which('tesseract')
    if w:
        return w
    for p in (r'C:\Program Files\Tesseract-OCR\tesseract.exe', r'C:\Program Files (x86)\Tesseract-OCR\tesseract.exe'):
        if Path(p).is_file():
            return p
    return None


def ocr_status():
    cmd = tesseract_cmd()
    if not cmd:
        return {'available': False, 'backend': None, 'langs': [],
                'message': 'Tesseract OCR is not installed. Install it (Windows: UB-Mannheim build) or set TESSERACT_CMD. '
                           'Without it you can still add/enter MIR lines manually next to the page image.'}
    try:
        out = subprocess.run([cmd, '--list-langs'], capture_output=True, text=True, timeout=30)
        langs = [x.strip() for x in (out.stdout + out.stderr).splitlines()[1:] if x.strip()]
    except Exception as e:
        return {'available': False, 'backend': 'tesseract', 'langs': [], 'message': 'Tesseract could not run: %s' % e}
    return {'available': True, 'backend': 'tesseract', 'cmd': cmd, 'langs': langs, 'message': 'Tesseract ready.'}


def render_pages(project_dir, doc, dpi=200):
    """Render every page to PNG (cached). Returns [{'page': n, 'path': Path}]."""
    rd = recon_dir(project_dir)
    pdf = stored_path(project_dir, 'mir', doc)
    outs = [(n, rd / 'pages' / ('%s_p%d.png' % (doc['id'], n))) for n in range(1, int(doc['pages']) + 1)]
    if all(p.is_file() for _, p in outs):
        return [{'page': n, 'path': p} for n, p in outs]
    try:
        import fitz
        with fitz.open(str(pdf)) as d:
            for n, p in outs:
                if not p.is_file():
                    d[n - 1].get_pixmap(matrix=fitz.Matrix(dpi / 72.0, dpi / 72.0)).save(str(p))
    except ImportError:
        tmp = Path(tempfile.mkdtemp(prefix='recon_'))
        try:
            subprocess.run(['pdftoppm', '-r', str(dpi), '-png', str(pdf), str(tmp / 'pg')], check=True, capture_output=True, timeout=300)
            files = sorted(tmp.glob('pg-*.png'))
            if len(files) != len(outs):
                raise ReconError('Page rendering produced %d images for %d pages.' % (len(files), len(outs)))
            for (n, p), f in zip(outs, files):
                shutil.move(str(f), str(p))
        except (OSError, subprocess.SubprocessError) as e:
            raise ReconError('Cannot render PDF pages (install PyMuPDF or poppler): %s' % e)
        finally:
            shutil.rmtree(tmp, ignore_errors=True)
    return [{'page': n, 'path': p} for n, p in outs]


def page_image_path(project_dir, st, doc_id, page):
    doc = next((d for d in st['mir_docs'] if d['id'] == doc_id), None)
    if not doc:
        raise ReconError('MIR document not found in this project.')
    page = int(page)
    if page < 1 or page > int(doc['pages']):
        raise ReconError('Page out of range.')
    return next(x['path'] for x in render_pages(project_dir, doc) if x['page'] == page)


def _tesseract_lines(img_path, langs=None, psm=6):
    """OCR one image (English only: table text is English; adding Arabic only injects garbage into descriptions)."""
    cmd = tesseract_cmd()
    out = subprocess.run([cmd, str(img_path), 'stdout', '--psm', str(psm), '-l', 'eng', 'tsv'],
                         capture_output=True, text=True, timeout=240, encoding='utf-8', errors='replace')
    if out.returncode != 0:
        raise ReconError('Tesseract failed: %s' % out.stderr.strip()[:200])
    lines = {}
    for row in out.stdout.splitlines()[1:]:
        p = row.split('\t')
        if len(p) < 12 or not p[11].strip():
            continue
        key = (p[2], p[3], p[4])
        try: conf = float(p[10])
        except ValueError: conf = -1
        e = lines.setdefault(key, {'words': [], 'confs': [], 'top': int(p[7])})
        e['words'].append(p[11].strip())
        if conf >= 0: e['confs'].append(conf)
    res = []
    for k in sorted(lines, key=lambda k: lines[k]['top']):
        e = lines[k]
        res.append({'text': ' '.join(e['words']), 'conf': round(sum(e['confs']) / len(e['confs']), 1) if e['confs'] else None})
    return res


def _image_variants(img_path, tmpdir):
    """Scans are photos with uneven light: no single preprocessing works for all, so we read 3 variants and vote."""
    try:
        import cv2
    except ImportError:
        return [('gray', str(img_path))]
    im = cv2.imread(str(img_path), cv2.IMREAD_GRAYSCALE)
    if im is None:
        return [('gray', str(img_path))]
    if im.shape[1] < 2400:
        sc = 2480.0 / im.shape[1]
        im = cv2.resize(im, None, fx=sc, fy=sc, interpolation=cv2.INTER_CUBIC)
    bg = cv2.medianBlur(cv2.resize(im, None, fx=.25, fy=.25), 31)
    bg = cv2.resize(bg, (im.shape[1], im.shape[0]))
    outs = [('gray', im), ('norm', cv2.divide(im, bg, scale=255)), ('clahe', cv2.createCLAHE(3.0, (8, 8)).apply(im))]
    res = []
    for name, img in outs:
        p = Path(tmpdir) / ('v_%s.png' % name)
        cv2.imwrite(str(p), img)
        res.append((name, str(p)))
    return res


_DATE_RE = re.compile(r'(?<!\d)(\d{1,2})\s*[/\-.]\s*(\d{1,2})\s*[/\-.]\s*(20[2-3]\d)(?!\d)')
_UNIT_WORDS = r'Rolls?|Pcs?|PC|Pieces?|Meters?|Mtr|M|Sets?|Box(?:es)?|Lot|Nos?|Each'
_ROW_ARABIC_FORM = re.compile(r'(?:^|[\s|])(?P<inv>\d{3,6})?[\s|]*(?P<qty>\d{1,5})\s*(?P<unit>%s)\b[\s|]*(?P<desc>[A-Za-z0-9][A-Za-z0-9 ()\-/.,#]{5,})' % _UNIT_WORDS, re.I)


def _clean_desc(s):
    s = re.split(r'[|\[\]{}]', s or '')
    s = next((x for x in s if sum(c.isalpha() for c in x) >= 3), '')
    s = re.sub(r'[^\x20-\x7e]', ' ', s)
    toks = [t for t in s.split() if re.fullmatch(r"[A-Za-z0-9()\-/.,#&+']+", t)]
    while toks and len(toks[0]) <= 2 and not re.search(r'\d', toks[0]) and toks[0] != '(':
        toks.pop(0)            # leading junk like "RE", "i)", "Le"
    while toks and len(toks[-1]) <= 2 and not toks[-1].isdigit() and not toks[-1].isupper():
        toks.pop()
    return ' '.join(toks).strip()


def _strip_lines(img_path, y0=0.12, y1=0.60, h=120, step=60):
    """Read the page in thin horizontal strips (one text line each). Slower, but copes with boxed/stamped forms
    that confuse Tesseract's page layout analysis."""
    try:
        import cv2
    except ImportError:
        return []
    im = cv2.imread(str(img_path), cv2.IMREAD_GRAYSCALE)
    if im is None:
        return []
    H = im.shape[0]
    cmd, out = tesseract_cmd(), []
    for y in range(int(H * y0), int(H * y1), step):
        strip = cv2.copyMakeBorder(im[y:y + h, :], 10, 10, 10, 10, cv2.BORDER_CONSTANT, value=255)
        tmp = Path(img_path).with_name('_strip.png')
        cv2.imwrite(str(tmp), strip)
        r = subprocess.run([cmd, str(tmp), 'stdout', '--psm', '7', '-l', 'eng'], capture_output=True, text=True, timeout=60,
                           encoding='utf-8', errors='replace')
        t = ' '.join(r.stdout.split())
        if len(re.sub(r'[^A-Za-z0-9]', '', t)) >= 3:
            out.append({'text': t, 'conf': None})
    return out


def _date_candidates(blob):
    out = []
    for d, m, y in _DATE_RE.findall(blob):
        if (1 <= int(d) <= 31 and 1 <= int(m) <= 12) or (1 <= int(m) <= 31 and 1 <= int(d) <= 12):
            out.append('%d/%d/%s' % (int(d), int(m), y))
    return out


def parse_mir_lines(lines):
    """One OCR pass -> {'ref_nos', 'dates', 'rows'} (all DRAFT, conservative)."""
    texts = [l['text'] for l in lines]
    blob = '\n'.join(texts)
    refs = []
    for m in re.finditer(r'MIR\s*[_#\-: ]*\s*([0-9OoIl]{2,4})\b', blob, re.I):
        digits = m.group(1).translate(str.maketrans('OoIl', '0011'))
        if digits.isdigit():
            refs.append('MIR-%03d' % int(digits))
    dates = _date_candidates(blob)
    start = next((i for i, t in enumerate(texts) if re.search(r'MATER\w*\s+DISC|MATER\w*\s+DESC|DISCRETION', t, re.I)), None)
    rows = []
    if start is not None:
        for i in range(start + 1, len(texts)):
            t = texts[i]
            if re.match(r'^\W*TOTAL\b', t, re.I) or re.search(r'PREPARED|PROJECT\s+MANAGER|CONSULT', t, re.I):
                break
            clean = re.sub(r'[|\[\]_—~=]+', ' ', t).strip()
            toks = [x for x in clean.split() if re.search(r'[A-Za-z0-9]', x)]   # drop stray "{" "-" etc.
            if len(toks) < 2:
                continue
            nums = []
            while toks and re.fullmatch(r'[\d,]+(\.\d+)?', toks[-1]) and len(nums) < 3:
                nums.insert(0, to_number(toks.pop()))
            unit, fz = (None, False)
            for _ in range(4):                         # skip OCR noise ("f", "s{") between the unit and the numbers
                if not toks:
                    break
                u, f = fuzzy_unit(toks[-1])
                if u:
                    unit, fz = u, f
                    toks.pop()
                    break
                if len(toks[-1]) <= 2 and not toks[-1].isdigit():
                    toks.pop()
                else:
                    break
            if toks and re.fullmatch(r'\d{1,2}', toks[0]):
                toks.pop(0)
            desc = _clean_desc(' '.join(toks))
            if not unit and not nums:
                if rows and len(desc) > 3:
                    rows[-1]['extra'] = (rows[-1].get('extra', '') + ' ' + desc).strip()
                continue
            if sum(c.isalpha() for c in desc) < 3:
                continue
            qty = total = None
            if len(nums) >= 2:
                qty, total = nums[-2], nums[-1]
            elif len(nums) == 1:
                total = nums[0]
            rows.append({'description': desc, 'unit': unit, 'unit_fuzzy': fz, 'qty': qty, 'total': total, 'conf': lines[i].get('conf')})
    if not rows:  # bilingual form: | invoice | 150 Roll | material name |
        for l in lines:
            m = _ROW_ARABIC_FORM.search(l['text'])
            if m:
                desc = _clean_desc(m.group('desc'))
                if sum(c.isalpha() for c in desc) >= 6:
                    u, fz = fuzzy_unit(m.group('unit'))
                    rows.append({'description': desc, 'unit': u, 'unit_fuzzy': fz, 'qty': to_number(m.group('qty')), 'total': None,
                                 'conf': l.get('conf'), 'invoice_no': m.group('inv')})
                    break
    return {'ref_nos': refs, 'dates': dates, 'rows': rows}


def _vote(values, min_votes=1):
    vals = [v for v in values if v not in (None, '')]
    if not vals:
        return None, []
    cnt = {}
    for v in vals:
        cnt[v] = cnt.get(v, 0) + 1
    ranked = sorted(cnt.items(), key=lambda kv: (-kv[1], vals.index(kv[0])))
    return (ranked[0][0] if ranked[0][1] >= min_votes else None), [(k, n) for k, n in ranked]


def merge_variant_parses(parses, n_variants, hint_ref=None):
    """Combine several OCR readings of the same page. Disagreements are kept as alternatives and flagged."""
    ref, _ = _vote([p for pr in parses for p in set(pr['ref_nos'])])
    date, _ = _vote([d for pr in parses for d in set(pr['dates'])])
    flags_page = []
    if not ref and hint_ref:
        ref = hint_ref; flags_page.append('ref_from_filename')
    first_rows = [pr['rows'][0] for pr in parses if pr['rows']]
    if not first_rows:
        return {'ref_no': ref, 'date': date, 'rows': [], 'flags': flags_page}
    desc_cands = [r['description'] for r in first_rows]
    tok_sets = [set(t.lower() for t in d.split()) for d in desc_cands]
    best = max(first_rows, key=lambda r: (r.get('conf') or 0) + 5 * len(r['description'].split()))
    toks = best['description'].split()
    if len(first_rows) >= 2:
        keep = [t for t in toks if sum(t.lower() in ts for ts in tok_sets) >= 2]
        if len(keep) >= 2:
            toks = keep
    desc = ' '.join(toks)
    unit, _ = _vote([r['unit'] for r in first_rows])
    fz = any(r['unit_fuzzy'] for r in first_rows if r['unit'] == unit)
    row = {'description': desc, 'unit': unit, 'conf': best.get('conf'), 'flags': [], 'alts': {}}
    for fld in ('qty', 'total'):
        v, ranked = _vote([r[fld] for r in first_rows])
        row[fld] = v
        if len(ranked) > 1:
            row['alts'][fld] = [k for k, _ in ranked]
            row['flags'].append('%s_variants_disagree' % fld)
        elif v is not None and len([r for r in first_rows if r[fld] == v]) < 2 and n_variants > 1:
            row['flags'].append('%s_single_reading' % fld)
    if fz: row['flags'].append('unit_fuzzy')
    if any(re.search(r'[a-z][A-Z]', t) or re.search(r'(.)\1\1', t) or (len(t) >= 6 and not re.search(r'[aeiouAEIOU0-9]', t)) for t in desc.split()):
        row['flags'].append('description_unreliable')
    if row['qty'] is None and row['total'] is not None: row['flags'].append('qty_blank_total_only')
    if row['qty'] is None and row['total'] is None: row['flags'].append('qty_not_read')
    return {'ref_no': ref, 'date': date, 'rows': [row], 'flags': flags_page}


def _new_record(doc, page, origin, ocr, flags=None):
    return {'id': uuid.uuid4().hex[:12], 'doc_id': doc['id'], 'page': page, 'origin': origin,
            'ocr': ocr, 'user': {}, 'review_status': 'pending', 'match': {'state': None, 'boq_id': None, 'candidates': []},
            'dup_override': None, 'ocr_flags': flags or [], 'created_at': _now(), 'edited_by': None, 'edited_at': None}


def _record_touched(r):
    return bool(r.get('user')) or r['review_status'] != 'pending' or (r.get('match') or {}).get('state') in ('approved', 'rejected')


def _ocr_page_legacy(pg, doc):
    """Fallback engine (no OpenCV, or the layout reader found nothing): 3 image variants read by Tesseract and voted."""
    tmp = tempfile.mkdtemp(prefix='recon_ocr_')
    try:
        variants = _image_variants(pg['path'], tmp)
        parses, raw = [], ''
        for name, p in variants:
            lines = _tesseract_lines(p)
            parses.append(parse_mir_lines(lines))
            if not raw:
                raw = '\n'.join(l['text'] for l in lines)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    if not any(p['rows'] for p in parses):
        tmp2 = tempfile.mkdtemp(prefix='recon_strip_')
        try:
            for name, p in _image_variants(pg['path'], tmp2)[:2]:
                parses.append(parse_mir_lines(_strip_lines(p)))
        finally:
            shutil.rmtree(tmp2, ignore_errors=True)
    merged = merge_variant_parses(parses, len(variants), doc.get('doc_ref_no_hint') or None)
    rows = [dict(r, row_flags=list(r.get('flags') or [])) for r in merged['rows']]
    return {'engine': 'variants', 'template': 'unknown', 'ref_no': merged['ref_no'], 'ref_source': 'header' if merged['ref_no'] else None,
            'date': merged['date'], 'date_alternatives': [], 'rows': rows, 'flags': list(merged['flags']),
            'status': {'value': None, 'confidence': 'none', 'ink': [], 'note': 'plain OCR: status not read'}, 'raw_text': raw[:4000]}


def _ocr_pages(project_dir, doc, status, progress=None):
    """Heavy part (touches no project state): render + read every page. Returns [(page_no, parsed_dict)]."""
    pages = render_pages(project_dir, doc)
    try:
        import recon_vision as V
        use_vision = V.available()
    except Exception:
        V, use_vision = None, False
    out = []
    for pg in pages:
        if progress:
            progress(pg['page'], len(pages))
        parsed = None
        if use_vision:
            try:
                prep, _ang = V.prepare(V.imread_gray(pg['path']))
                parsed = V.read_page(status['cmd'], prep, doc.get('doc_ref_no_hint'), 'eng', fuzzy_unit)
                parsed['engine'] = 'layout'
            except Exception as e:  # never fail the whole PDF because one reader broke
                print('[RECON] layout reader failed on %s p%d (%s); using variant voting' % (doc.get('filename'), pg['page'], e))
                parsed = None
        if parsed is None or not parsed.get('rows'):
            legacy = _ocr_page_legacy(pg, doc)
            if parsed is None:
                parsed = legacy
            elif legacy['rows']:  # layout reader read the header/status but not the table: keep its header, take the legacy rows
                parsed['rows'] = legacy['rows']
                parsed['flags'] = [f for f in parsed.get('flags', []) if f != 'table_header_not_found'] + ['rows_from_variant_vote']
                parsed['engine'] = 'layout+variants'
        out.append((pg['page'], parsed))
    return out


def _supplier_from_text(st, text):
    """The old form prints the supplier ("Commscope") in its own field, outside the item description. If a word on the page
    is a BOQ brand (typos tolerated) it is kept as evidence for matching only; it never changes quantities."""
    keys = {brand_key(i['brand']): i['brand'] for i in (st.get('boq') or {}).get('items', []) if brand_key(i['brand'])}
    for w in re.findall(r'[A-Za-z]{5,}', text or ''):
        k = brand_key(w)
        for bk, name in keys.items():
            if k == bk or (len(bk) >= 6 and difflib.SequenceMatcher(None, norm_alnum(w), bk).ratio() >= 0.85):
                return name
    return None


def _apply_ocr(st, doc, results, user):
    """Light part: turn page results into DRAFT records (OCR values are kept untouched in rec['ocr'])."""
    n_new = 0
    for page_no, parsed in results:
        stat = parsed.get('status') or {}
        base = {'ref_no': parsed.get('ref_no'), 'date': parsed.get('date'), 'date_alternatives': parsed.get('date_alternatives') or [],
                'ref_source': parsed.get('ref_source'), 'template': parsed.get('template'), 'engine': parsed.get('engine'),
                'supplier': parsed.get('supplier'), 'mir_status': stat.get('value'), 'status_note': stat.get('note'),
                'status_ink': stat.get('ink') or []}
        pflags = list(parsed.get('flags') or [])
        base['supplier'] = base.get('supplier') or _supplier_from_text(st, parsed.get('raw_text') or '')
        if parsed.get('template') in ('new', 'old') and not stat.get('value'):
            pflags.append('status_unclear')
        rows = parsed.get('rows') or []
        if not rows:
            rec = _new_record(doc, page_no, 'ocr', dict(base, description=None, unit=None, qty=None, total=None, conf=None,
                                                         raw_text=parsed.get('raw_text', '')), list(dict.fromkeys(pflags + ['ocr_no_lines'])))
            st['records'].append(rec); n_new += 1
            continue
        for row in rows:
            fl = list(pflags) + list(row.get('row_flags') or [])
            if row.get('unit_fuzzy') and 'old_form_partial' not in fl and 'unit_fuzzy' not in fl:
                fl.append('unit_fuzzy')
            if row.get('qty') is None and row.get('total') is not None and 'qty_blank_total_only' not in fl: fl.append('qty_blank_total_only')
            if row.get('qty') is None and row.get('total') is None and 'qty_not_read' not in fl: fl.append('qty_not_read')
            # text below the item line (often handwriting) is kept for the reviewer but never mixed into the description used for matching
            rec = _new_record(doc, page_no, 'ocr', dict(base, description=row.get('description') or None, unit=row.get('unit'), qty=row.get('qty'),
                                                         total=row.get('total'), conf=row.get('conf'), invoice_no=row.get('invoice_no'),
                                                         alts=row.get('alts'), extra_text=row.get('extra') or None, raw_text=''), list(dict.fromkeys(fl)))
            st['records'].append(rec); n_new += 1
    doc.update({'ocr_status': 'done', 'ocr_backend': 'tesseract (layout reader + 3-variant vote)', 'ocr_at': _now(), 'ocr_error': None, 'ocr_progress': None})
    _audit(st, user, 'MIR_OCR', '%s records=%d' % (doc['filename'], n_new))
    return n_new


def _check_ocr_allowed(st, doc_id):
    doc = next((d for d in st['mir_docs'] if d['id'] == doc_id), None)
    if not doc:
        raise ReconError('MIR document not found in this project.')
    status = ocr_status()
    if not status['available']:
        raise ReconError(status['message'])
    if any(_record_touched(r) for r in st['records'] if r['doc_id'] == doc_id):
        raise ReconError('This document already has edited/reviewed records. Delete them first to re-run OCR (prevents double counting).')
    return doc, status


def run_ocr(project_dir, st, doc_id, user):
    """Synchronous OCR (tests / scripts). The web UI uses start_ocr()."""
    doc, status = _check_ocr_allowed(st, doc_id)
    results = _ocr_pages(project_dir, doc, status)
    st['records'] = [r for r in st['records'] if r['doc_id'] != doc_id]
    return _apply_ocr(st, doc, results, user)


def start_ocr(project_dir, st, doc_id, user):
    """Start OCR in a background thread so the web server is not blocked (a page takes ~10-20 s).
    `st` is the caller's already-locked state: the 'running' marker is stored by the caller's save."""
    doc, status = _check_ocr_allowed(st, doc_id)
    if doc.get('ocr_status') == 'running' and doc.get('ocr_proc') == _PROC_ID:
        raise ReconError('OCR is already running for this PDF.')
    doc.update({'ocr_status': 'running', 'ocr_proc': _PROC_ID, 'ocr_started_at': _now(), 'ocr_error': None,
                'ocr_progress': {'page': 0, 'pages': doc['pages']}})
    _audit(st, user, 'MIR_OCR_START', doc['filename'])
    threading.Thread(target=_ocr_worker, args=(str(project_dir), dict(doc), status, dict(user or {})), daemon=True, name='recon-ocr').start()
    return True


def _ocr_worker(project_dir, doc, status, user):
    def set_doc(fn):
        with locked_state(project_dir) as st:
            d = next((x for x in st['mir_docs'] if x['id'] == doc['id']), None)
            if d is not None:
                fn(st, d)
            return d is not None

    def progress(page, pages):
        set_doc(lambda st, d: d.__setitem__('ocr_progress', {'page': page, 'pages': pages}))
    try:
        with _plock(project_dir):  # wait until the caller's save (which marks the job 'running') has been written
            if not (Path(project_dir) / 'reconciliation' / 'state.json').exists():
                return  # project was deleted meanwhile: never recreate its folder
        results = _ocr_pages(project_dir, doc, status, progress)

        def finish(st, d):
            if any(_record_touched(r) for r in st['records'] if r['doc_id'] == d['id']):
                raise ReconError('Records of this PDF were edited while OCR was running; OCR result discarded.')
            st['records'] = [r for r in st['records'] if r['doc_id'] != d['id']]
            _apply_ocr(st, d, results, user)
        set_doc(finish)
    except Exception as e:
        msg = str(e)[:300] or e.__class__.__name__
        print('[RECON] OCR failed for %s: %s' % (doc.get('filename'), msg))
        try:
            set_doc(lambda st, d: d.update({'ocr_status': 'error', 'ocr_error': msg, 'ocr_progress': None}))
        except Exception:
            pass


# ----------------------------------------------------------------------------
# records: manual add / edit / review
# ----------------------------------------------------------------------------
EDITABLE = ('description', 'unit', 'qty', 'total', 'ref_no', 'date', 'mir_status', 'brand', 'note')


def add_manual_record(st, doc_id, page, user):
    doc = next((d for d in st['mir_docs'] if d['id'] == doc_id), None)
    if not doc:
        raise ReconError('MIR document not found in this project.')
    page = int(page or 1)
    if page < 1 or page > int(doc['pages']):
        raise ReconError('Page out of range.')
    rec = _new_record(doc, page, 'manual', {'description': None, 'unit': None, 'qty': None, 'total': None,
                                            'ref_no': None, 'date': None, 'conf': None, 'raw_text': ''}, ['manual_entry'])
    rec['edited_by'], rec['edited_at'] = (user or {}).get('username', ''), _now()
    st['records'].append(rec)
    _audit(st, user, 'MIR_RECORD_ADD', '%s p%d' % (doc['filename'], page))
    return rec


def _clean_field(k, v):
    if v is None:
        return None
    if k in ('qty', 'total'):
        if str(v).strip() == '':
            return None
        n = to_number(v)
        if n is None:
            raise ReconError('"%s" must be a number (got "%s").' % (k, v))
        if n < 0:
            raise ReconError('"%s" cannot be negative.' % k)
        return n
    if k == 'mir_status':
        s = str(v).strip().upper()
        if s == '':
            return None
        if s not in MIR_STATUS_LABELS:
            raise ReconError('MIR status must be A, B, C or D.')
        return s
    return _txt(v) or None


def eff(rec, key):
    u = rec.get('user', {})
    return u[key] if key in u else rec.get('ocr', {}).get(key)


def update_record(st, rec_id, fields, mark_reviewed, user):
    r = next((x for x in st['records'] if x['id'] == rec_id), None)
    if not r:
        raise ReconError('Record not found.')
    changed = []
    for k, v in (fields or {}).items():
        if k not in EDITABLE:
            continue
        nv = _clean_field(k, v)
        if r['user'].get(k, '__unset__') != nv:
            changed.append(k)
        r['user'][k] = nv
    if mark_reviewed:
        # snapshot every effective value so what was approved is exactly what is stored
        for k in EDITABLE:
            if k not in r['user']:
                r['user'][k] = _clean_field(k, r['ocr'].get(k)) if k in r['ocr'] else None
        missing = []
        if not r['user'].get('description'): missing.append('description')
        if not r['user'].get('unit'): missing.append('unit')
        if r['user'].get('qty') is None: missing.append('qty (Q.T)')
        if not r['user'].get('mir_status'): missing.append('MIR status (A/B/C/D)')
        if not r['user'].get('ref_no'): missing.append('MIR ref no.')
        if missing:
            raise ReconError('Cannot mark as reviewed - missing: %s. (Values are never guessed.)' % ', '.join(missing))
        r['review_status'] = 'reviewed'
    elif changed:
        r['review_status'] = 'pending'  # an edit invalidates a previous review
    r['edited_by'], r['edited_at'] = (user or {}).get('username', ''), _now()
    _audit(st, user, 'MIR_RECORD_REVIEW' if mark_reviewed else 'MIR_RECORD_EDIT', '%s fields=%s' % (rec_id, ','.join(changed) or '-'))
    return r


def delete_record(st, rec_id, user):
    r = next((x for x in st['records'] if x['id'] == rec_id), None)
    if not r:
        raise ReconError('Record not found.')
    st['records'].remove(r)
    _audit(st, user, 'MIR_RECORD_DELETE', '%s origin=%s' % (rec_id, r['origin']))


def delete_doc(project_dir, st, doc_id, user):
    doc = next((d for d in st['mir_docs'] if d['id'] == doc_id), None)
    if not doc:
        raise ReconError('MIR document not found.')
    st['records'] = [r for r in st['records'] if r['doc_id'] != doc_id]
    st['mir_docs'].remove(doc)
    # the original file is kept on disk (evidence); it is only removed from the reconciliation set
    st.setdefault('removed_docs', []).append({**doc, 'removed_at': _now(), 'removed_by': (user or {}).get('username', '')})
    _audit(st, user, 'MIR_DOC_REMOVE', doc['filename'])


# ----------------------------------------------------------------------------
# matching
# ----------------------------------------------------------------------------
_KINDS = (('patch cord', r'patch|cord|jumper'), ('faceplate', r'face\s*plate|faceplate'), ('outlet', r'outlet|jack|keystone|socket'),
          ('panel', r'panel'), ('organizer', r'organi[sz]er|manager'), ('cable', r'cables?\b|utp|ftp|stp'))


def _kind(text):
    t = (text or '').lower()
    for name, rx in _KINDS:
        if re.search(rx, t):
            return name
    return None


def _score_candidates(st, rec):
    items = (st.get('boq') or {}).get('items', [])
    desc = eff(rec, 'description') or ''
    mir_unit = canon_unit(eff(rec, 'unit'))
    mir_brand = eff(rec, 'brand') or (rec.get('ocr') or {}).get('supplier') or ''
    d_alnum = norm_alnum(desc)
    d_tok = tokens(desc)
    d_models = model_tokens(desc)
    rejected = set((rec.get('match') or {}).get('rejected_ids', []))
    out = []
    for it in items:
        if it['id'] in rejected:
            continue
        reasons, score, strong = [], 0.0, None
        part_n = norm_alnum(it['part'])
        ref_n = norm_alnum(it['reference'])
        if len(part_n) >= 5 and part_n in d_alnum:
            score += 0.60; reasons.append('part # %s found in MIR text' % it['part']); strong = ('p', part_n)
        elif len(ref_n) >= 6 and ref_n in d_alnum.replace('accessories', ''):
            score += 0.60; reasons.append('BOQ reference text found in MIR text'); strong = ('r', ref_n)
        else:
            bm = model_tokens(it['part'] + ' ' + it['reference'] + ' ' + it['item'])
            common = d_models & bm
            if common:
                score += 0.35; reasons.append('model %s' % ', '.join(sorted(common))); strong = ('r', ref_n or part_n or it['id'])
        bk = brand_key(it['brand'])
        if bk:
            bt = [norm_alnum(t) for t in re.findall(r'[A-Za-z]+', desc + ' ' + mir_brand)]
            if brand_key(mir_brand) == bk or any(brand_key(t) == bk or (len(t) >= 6 and len(bk) >= 6 and difflib.SequenceMatcher(None, t, bk).ratio() >= 0.85) for t in bt):
                score += 0.15; reasons.append('brand %s' % it['brand'])
        pool = tokens(it['item'] + ' ' + it['reference'] + ' ' + it['brand'])
        label = tokens(it['item'])
        if d_tok and pool:
            cov = _hits(d_tok, pool) / len(d_tok)
            score += 0.25 * cov
            if cov: reasons.append('%d%% of MIR words found in BOQ text' % round(cov * 100))
        if label and d_tok:
            bov = _hits(label, d_tok) / len(label)
            score += 0.15 * bov
        kind_it = _kind(it['item'] + ' ' + it['reference'])
        kind_mir = _kind(desc)
        if kind_it and kind_mir and kind_it != kind_mir:
            score -= 0.15; reasons.append('different kind of material (%s vs %s)' % (kind_mir, kind_it))
        ub, _src = boq_unit(st, it)
        ustate = 'boq_unit_unknown' if not ub else ('mir_unit_missing' if not mir_unit else ('ok' if ub == mir_unit else 'mismatch'))
        if ustate == 'ok':
            score += 0.05; reasons.append('unit %s' % mir_unit)
        if score > 0:
            out.append({'boq_id': it['id'], 'score': round(min(score, 1.0), 3), 'reasons': reasons, 'unit_state': ustate,
                        'strong': list(strong) if strong else None, 'exact_ev': bool(strong and (strong[0] == 'p' or reasons and 'reference' in ' '.join(reasons)))})
    out.sort(key=lambda c: -c['score'])
    return out


def classify_candidates(cands):
    if not cands:
        return 'unmatched', None
    top = cands[0]
    if top['score'] < 0.30:
        # weak but clearly leading: offered as a suggestion only (never auto-approved); the reviewer decides
        if top['score'] >= 0.18 and (len(cands) == 1 or top['score'] - cands[1]['score'] >= 0.05):
            return 'suggested', top['boq_id']
        return 'unmatched', None
    near = [c for c in cands[1:] if top['score'] - c['score'] <= 0.03]
    same_ev = [c for c in cands[1:] if top['strong'] and c['strong'] and c['strong'] == top['strong']]
    if near or same_ev:
        return 'ambiguous', None
    if top['exact_ev'] and top['score'] >= 0.60:
        return 'exact', top['boq_id']
    return 'suggested', top['boq_id']


def run_matching(st, user):
    n = {'exact': 0, 'suggested': 0, 'ambiguous': 0, 'unmatched': 0, 'kept': 0}
    for r in st['records']:
        m = r.get('match') or {}
        if m.get('state') == 'approved':
            n['kept'] += 1
            continue
        if not eff(r, 'description'):
            r['match'] = {'state': 'unmatched', 'boq_id': None, 'candidates': [], 'method': 'auto', 'rejected_ids': m.get('rejected_ids', []),
                          'note': 'No description to match - enter it in review.'}
            n['unmatched'] += 1
            continue
        cands = _score_candidates(st, r)
        state, bid = classify_candidates(cands)
        r['match'] = {'state': state, 'boq_id': bid, 'candidates': cands[:6], 'method': 'auto',
                      'rejected_ids': m.get('rejected_ids', []), 'decided_by': None, 'decided_at': None}
        n[state] += 1
    _audit(st, user, 'MATCH_RUN', json.dumps(n))
    return n


def decide_match(st, rec_id, action, boq_id, user):
    r = next((x for x in st['records'] if x['id'] == rec_id), None)
    if not r:
        raise ReconError('Record not found.')
    idx = boq_index(st)
    m = r.get('match') or {}
    who = (user or {}).get('username', '')
    if action == 'approve':
        target = boq_id or m.get('boq_id')
        if not target:
            raise ReconError('Choose the BOQ row to match (ambiguous/unmatched records need an explicit choice).')
        if target not in idx:
            raise ReconError('BOQ row not found.')
        manual = (target != m.get('boq_id')) or m.get('state') in ('ambiguous', 'unmatched', 'needs_rematch', None)
        r['match'] = {**m, 'state': 'approved', 'boq_id': target, 'method': 'manual' if manual else 'auto_approved',
                      'decided_by': who, 'decided_at': _now()}
    elif action == 'reject':
        rej = list(m.get('rejected_ids', []))
        t = boq_id or m.get('boq_id')
        if t and t not in rej: rej.append(t)
        r['match'] = {'state': 'rejected', 'boq_id': None, 'candidates': [], 'method': 'manual', 'rejected_ids': rej,
                      'decided_by': who, 'decided_at': _now()}
    elif action == 'clear':
        r['match'] = {'state': None, 'boq_id': None, 'candidates': [], 'rejected_ids': m.get('rejected_ids', [])}
    else:
        raise ReconError('Unknown match action.')
    _audit(st, user, 'MATCH_' + action.upper(), '%s -> %s' % (rec_id, r['match'].get('boq_id') or '-'))
    return r


def bulk_approve_exact(st, user):
    n = 0
    for r in st['records']:
        m = r.get('match') or {}
        if m.get('state') == 'exact' and m.get('boq_id'):
            decide_match(st, r['id'], 'approve', m['boq_id'], user); n += 1
    return n


def set_boq_override(st, item_id, unit, note, user):
    if item_id not in boq_index(st):
        raise ReconError('BOQ row not found.')
    ov = st['boq_overrides'].setdefault(item_id, {})
    if unit is not None:
        ov['unit'] = _txt(unit) or None
    if note is not None:
        ov['note'] = _txt(note) or None
    ov['by'], ov['at'] = (user or {}).get('username', ''), _now()
    _audit(st, user, 'BOQ_ITEM_OVERRIDE', '%s unit=%s' % (item_id, ov.get('unit')))


# ----------------------------------------------------------------------------
# calculation (pure function of state; nothing here mutates stored data)
# ----------------------------------------------------------------------------
def _r(v, n=4):
    return None if v is None else round(v, n)


def compute(st):
    boq = st.get('boq') or {}
    items = boq.get('items', [])
    idx = {it['id']: it for it in items}
    count_status = set(st['settings'].get('count_statuses') or COUNT_STATUSES_DEFAULT)
    docs = {d['id']: d for d in st['mir_docs']}
    exceptions = []

    def exc(sev, typ, ref, msg):
        exceptions.append({'severity': sev, 'type': typ, 'ref': ref, 'message': msg})

    # Make document/page coverage explicit. A page with no extracted line may be
    # blank or may have been missed by OCR; never silently assume either.
    records_by_doc_page = {}
    for record in st.get('records', []):
        records_by_doc_page.setdefault((record.get('doc_id'), int(record.get('page') or 1)), 0)
        records_by_doc_page[(record.get('doc_id'), int(record.get('page') or 1))] += 1
    for doc in st.get('mir_docs', []):
        doc_id = doc.get('id')
        filename = doc.get('filename') or doc_id or '(unknown MIR file)'
        ocr_status = doc.get('ocr_status') or 'not_run'
        if ocr_status != 'done':
            severity = 'error' if ocr_status == 'error' else 'warning'
            detail = ('OCR failed: %s' % (doc.get('ocr_error') or 'no error details')) if ocr_status == 'error' else ('OCR status is %s' % ocr_status)
            exc(severity, 'mir_document_not_fully_read', str(filename), '%s: %s. Verify every page manually; unread documents are not assumed empty.' % (filename, detail))
        try:
            page_count = max(0, int(doc.get('pages') or 0))
        except (TypeError, ValueError):
            page_count = 0
        for page_no in range(1, page_count + 1):
            if not records_by_doc_page.get((doc_id, page_no), 0):
                exc('warning', 'mir_page_without_records', '%s p%d' % (filename, page_no),
                    '%s page %d has no extracted or manually added MIR lines. Check the scan; if a line was missed, add it manually.' % (filename, page_no))

    for w in boq.get('warnings', []):
        exc(w.get('severity', 'info'), 'boq_' + w.get('type', ''), w.get('ref', ''), w.get('message', ''))

    # ---- record views + flags
    ordered = sorted(st['records'], key=lambda r: (docs.get(r['doc_id'], {}).get('uploaded_at', ''), r['doc_id'], r['page'], r['created_at']))
    keyseen, recs = {}, []
    ref_docs = {}
    for r in ordered:
        d = docs.get(r['doc_id'])
        if not d:
            continue
        e = {k: eff(r, k) for k in EDITABLE}
        ref_digits = re.sub(r'\D', '', e['ref_no'] or '')
        if ref_digits:
            ref_docs.setdefault(ref_digits.lstrip('0') or '0', {}).setdefault(r['doc_id'], set()).add(norm_alnum(e['description']))
        flags = list(r.get('ocr_flags', []))
        if r['review_status'] != 'reviewed': flags.append('not_reviewed')
        if e['qty'] is None: flags.append('qty_missing')
        if e['qty'] is not None and e['total'] is not None and abs(e['qty'] - e['total']) > 1e-9: flags.append('qty_total_differ')
        if not e['mir_status']: flags.append('status_missing')
        elif e['mir_status'] not in count_status: flags.append('status_not_counted_' + e['mir_status'])
        if not e['unit']: flags.append('unit_missing')
        if not e['ref_no']: flags.append('ref_no_missing')
        c = (r.get('ocr') or {}).get('conf')
        if r['origin'] == 'ocr' and c is not None and c < 60: flags.append('low_ocr_confidence')
        if 'mir_status' not in (r.get('user') or {}) and (r.get('ocr') or {}).get('mir_status'):
            flags.append('status_from_ocr')  # read from the check-box ink by OCR; the reviewer must confirm it on the page
        dup_of = None
        if r.get('dup_override') == 'duplicate':
            dup_of = 'REVIEWER'; flags.append('duplicate')  # the reviewer decided this line must not be counted twice
        elif e['description'] and e['qty'] is not None and r.get('dup_override') != 'not_duplicate':
            key = (ref_digits, norm_alnum(e['description']), canon_unit(e['unit']), e['qty'], (e['date'] or ''))
            if key in keyseen:
                dup_of = keyseen[key]; flags.append('duplicate')
            else:
                keyseen[key] = r['id']
        m = r.get('match') or {}
        bid = m.get('boq_id') if m.get('state') == 'approved' else None
        recs.append({'id': r['id'], 'doc_id': r['doc_id'], 'filename': d['filename'], 'page': r['page'], 'origin': r['origin'],
                     'ocr': r['ocr'], 'user': r['user'], 'eff': e, 'review_status': r['review_status'], 'flags': flags,
                     'dup_of': dup_of, 'match': m, 'approved_boq_id': bid, 'edited_by': r.get('edited_by'), 'edited_at': r.get('edited_at')})
    groups = []  # same item + quantity on different MIR numbers: warn, flag, never auto-exclude
    for rv in recs:
        e = rv['eff']
        if not (e['description'] and e['qty'] is not None):
            continue
        for g in groups:
            g0 = g[0]['eff']
            if canon_unit(g0['unit']) == canon_unit(e['unit']) and g0['qty'] == e['qty'] and _similar_desc(g0['description'], e['description']):
                g.append(rv); break
        else:
            groups.append([rv])
    for g in groups:
        refs = {re.sub(r'\D', '', x['eff']['ref_no'] or '').lstrip('0') or '?' for x in g}
        if len(g) > 1 and len(refs) > 1 and len({x['doc_id'] for x in g}) > 1:
            e0 = g[0]['eff']
            for x in g:
                if 'duplicate' not in x['flags']:
                    x['flags'].append('possible_resubmission')
            exc('warning', 'possible_resubmission', ','.join(x['id'] for x in g),
                '%s carry the same item and quantity (%g %s of "%s"). Confirm whether these are separate deliveries or the same material '
                're-submitted; use "Exclude as duplicate" on any line that must not be counted twice. Nothing was excluded automatically.' % (
                    ', '.join('%s (%s%s)' % (x['eff']['ref_no'] or 'no MIR no.', x['filename'], ', status ' + x['eff']['mir_status'] if x['eff']['mir_status'] else '') for x in g),
                    e0['qty'], (canon_unit(e0['unit']) or ''), (e0['description'] or '')[:60]))
    for num, per_doc in ref_docs.items():
        if len(per_doc) > 1:
            names = ', '.join(docs[i]['filename'] for i in per_doc)
            descs = set().union(*per_doc.values())
            exc('warning', 'ref_no_conflict', 'MIR-%s' % num.zfill(3),
                'MIR number %s appears on %d different PDFs (%s)%s.' % (num, len(per_doc), names,
                    ' with different contents - confirm they are distinct requests' if len(descs) > 1 else ' with the same content - possible duplicate upload'))
            for rv in recs:
                if re.sub(r'\D', '', rv['eff']['ref_no'] or '').lstrip('0') == num and rv['doc_id'] in per_doc:
                    rv['flags'].append('ref_no_conflict')

    # ---- per BOQ item
    by_item = {it['id']: [] for it in items}
    for rv in recs:
        if rv['approved_boq_id'] in by_item:
            by_item[rv['approved_boq_id']].append(rv)
    pending_ptr = {it['id']: 0 for it in items}
    for rv in recs:
        m = rv['match']
        if m.get('state') in ('suggested', 'exact') and m.get('boq_id') in pending_ptr:
            pending_ptr[m['boq_id']] += 1
    rows = []
    for it in items:
        ub, usrc = boq_unit(st, it)
        qty, up = it['qty'], it['unit_price']
        flags = list(it['flags'])
        recorded = confirmed = pending = notappr = uconf = 0.0
        by_status = {'A': 0.0, 'B': 0.0, 'C': 0.0, 'D': 0.0, 'unknown': 0.0}
        sources, awaiting = [], pending_ptr[it['id']]
        unit_unverified = False
        for rv in by_item[it['id']]:
            e = rv['eff']
            q = e['qty']
            if 'duplicate' in rv['flags']:
                flags.append('duplicate_record_excluded'); continue
            if q is None:
                awaiting += 1; continue
            recorded += q
            st_ = e['mir_status'] if e['mir_status'] in by_status else 'unknown'
            mu = canon_unit(e['unit'])
            conflict = bool(ub and mu and ub != mu)
            if rv['review_status'] != 'reviewed' or st_ == 'unknown':
                pending += q; awaiting += 1; by_status['unknown'] += q if st_ == 'unknown' else 0
            elif conflict:
                uconf += q; flags.append('unit_conflict')
            elif st_ in count_status:
                confirmed += q; by_status[st_] += q
                if not ub: unit_unverified = True
            else:
                notappr += q; by_status[st_] += q
            sources.append({'file': rv['filename'], 'page': rv['page'], 'ref_no': e['ref_no'], 'date': e['date'], 'qty': q,
                            'unit': e['unit'], 'status': e['mir_status'], 'reviewed': rv['review_status'] == 'reviewed', 'record_id': rv['id']})
        unavailable = {}
        qty_ok = qty is not None and qty >= 0
        boq_value = qty * up if (qty_ok and up is not None and up >= 0) else None
        if not qty_ok: unavailable['boq_qty'] = 'BOQ quantity missing/invalid'
        if up is None or up < 0: unavailable['price'] = 'unit price missing/invalid - no financial values'
        remaining = (qty - confirmed) if qty_ok else None
        qty_pct = (confirmed / qty * 100) if (qty_ok and qty > 0) else None
        if qty_pct is None: unavailable['qty_pct'] = 'BOQ quantity is zero or unavailable'
        mir_value = confirmed * up if (up is not None and up >= 0) else None
        fin_pct = (mir_value / boq_value * 100) if (boq_value is not None and boq_value > 0 and mir_value is not None) else None
        if fin_pct is None and 'price' not in unavailable: unavailable['fin_pct'] = 'BOQ value is zero or unavailable'
        rem_value = (boq_value - mir_value) if (boq_value is not None and mir_value is not None) else None
        over = qty_ok and confirmed > qty + 1e-9
        if over:
            flags.append('over_boq')
            exc('error', 'over_boq', it['id'], '%s: confirmed MIR qty %g exceeds BOQ qty %g (remaining %g). Not capped.' % (it['id'], confirmed, qty, qty - confirmed))
        if unit_unverified: flags.append('unit_unverified')
        if uconf: exc('warning', 'unit_conflict', it['id'], '%s: %g MIR units excluded - MIR unit differs from BOQ unit %s.' % (it['id'], uconf, ub))
        if 'price_missing' in flags: exc('warning', 'price_missing', it['id'], '%s: no unit price; financial values unavailable.' % it['id'])
        if notappr: exc('info', 'status_not_counted', it['id'], '%s: %g recorded under MIR status C/D, not counted.' % (it['id'], notappr))
        if over: status = 'over_boq'
        elif qty_ok and qty > 0 and abs(confirmed - qty) < 1e-9: status = 'complete'
        elif confirmed > 0: status = 'partial'
        elif awaiting or pending: status = 'pending_review'
        elif by_item[it['id']]: status = 'no_confirmed'
        else: status = 'no_mir'
        rows.append({'item_id': it['id'], 'code': it['sn'] or '(no S.N)', 'row_ref': it['id'], 'description': it['item'] or it['reference'],
                     'reference': it['reference'], 'part': it['part'], 'brand': it['brand'] or '(no brand)', 'brand_key': brand_key(it['brand']) or '?',
                     'unit': ub, 'unit_source': usrc, 'boq_qty': qty, 'unit_price': up, 'boq_value': _r(boq_value, 2), 'sheet_total': it['total_price'],
                     'recorded_qty': _r(recorded), 'confirmed_qty': _r(confirmed), 'pending_qty': _r(pending), 'not_approved_qty': _r(notappr),
                     'unit_conflict_qty': _r(uconf), 'remaining_qty': _r(remaining), 'qty_pct': _r(qty_pct, 2), 'mir_value': _r(mir_value, 2),
                     'remaining_value': _r(rem_value, 2), 'fin_pct': _r(fin_pct, 2), 'unavailable': unavailable, 'status': status,
                     'flags': sorted(set(flags)), 'sources': sources, 'by_status': {k: _r(v) for k, v in by_status.items()},
                     'awaiting_review': awaiting, 'n_records': len(by_item[it['id']]), 'note': st['boq_overrides'].get(it['id'], {}).get('note')})

    # ---- unmatched MIR records
    names = {brand_key(it['brand']): it['brand'] for it in items if it['brand']}
    for rv in recs:
        if rv['approved_boq_id'] is None:
            d = (rv['eff']['description'] or '') + ' ' + (rv['eff']['brand'] or '')
            bk = brand_key(rv['eff']['brand']) if rv['eff']['brand'] else None
            if not bk:
                for k, nm in names.items():
                    if any(difflib.SequenceMatcher(None, norm_alnum(t), k).ratio() >= 0.85 for t in re.findall(r'[A-Za-z]{5,}', d)):
                        bk = k; break
            rv['brand_key'] = bk or '?'
            sm = rv['match'].get('state') or 'not_matched'
            exc('warning', 'mir_' + sm, rv['id'], '%s p%d: MIR line "%s" (%s) is %s - not counted.' % (rv['filename'], rv['page'], rv['eff']['description'] or '(blank)', rv['eff']['qty'], sm.replace('_', ' ')))
        if 'duplicate' in rv['flags']:
            if rv['dup_of'] == 'REVIEWER':
                exc('info', 'duplicate_record', rv['id'], '%s p%d (%s %s %s) was marked as a duplicate by the reviewer - excluded from totals.' % (rv['filename'], rv['page'], rv['eff']['ref_no'] or '?', rv['eff']['qty'], rv['eff']['unit'] or ''))
            else:
                exc('warning', 'duplicate_record', rv['id'], '%s p%d duplicates record %s - excluded until you mark it "not a duplicate".' % (rv['filename'], rv['page'], rv['dup_of']))

    # ---- brands (no mixing of units; financials only where valid)
    brands = {}
    for row in rows:
        b = brands.setdefault(row['brand_key'], {'brand': row['brand'], 'key': row['brand_key'], 'items': 0, 'items_valued': 0, 'items_unvalued': 0,
                                                 'qty_by_unit': {}, 'unspecified_unit_items': 0, 'boq_value': 0.0, 'mir_value': 0.0, 'remaining_value': 0.0,
                                                 'items_with_confirmed': 0, 'missing_mir': [], 'over_boq': [], 'awaiting_review': [], 'unmatched_mir_records': 0})
        b['items'] += 1
        if row['boq_value'] is not None and row['mir_value'] is not None:
            b['items_valued'] += 1; b['boq_value'] += row['boq_value']; b['mir_value'] += row['mir_value']; b['remaining_value'] += row['remaining_value']
        else:
            b['items_unvalued'] += 1
        if row['unit'] and row['boq_qty'] is not None:
            u = b['qty_by_unit'].setdefault(row['unit'], {'boq': 0.0, 'confirmed': 0.0})
            u['boq'] += row['boq_qty']; u['confirmed'] += row['confirmed_qty'] or 0
        else:
            b['unspecified_unit_items'] += 1
        if (row['confirmed_qty'] or 0) > 0: b['items_with_confirmed'] += 1
        if row['status'] == 'no_mir': b['missing_mir'].append(row['row_ref'])
        if row['status'] == 'over_boq': b['over_boq'].append(row['row_ref'])
        if row['awaiting_review']: b['awaiting_review'].append(row['row_ref'])
    for rv in recs:
        if rv['approved_boq_id'] is None:
            k = rv.get('brand_key', '?')
            brands.setdefault(k, {'brand': '(unknown brand)' if k == '?' else names.get(k, k), 'key': k, 'items': 0, 'items_valued': 0, 'items_unvalued': 0,
                                  'qty_by_unit': {}, 'unspecified_unit_items': 0, 'boq_value': 0.0, 'mir_value': 0.0, 'remaining_value': 0.0,
                                  'items_with_confirmed': 0, 'missing_mir': [], 'over_boq': [], 'awaiting_review': [], 'unmatched_mir_records': 0})['unmatched_mir_records'] += 1
    brand_list = []
    for b in brands.values():
        for u in b['qty_by_unit'].values():
            u['remaining'] = _r(u['boq'] - u['confirmed']); u['pct'] = _r(u['confirmed'] / u['boq'] * 100, 2) if u['boq'] > 0 else None
            u['boq'], u['confirmed'] = _r(u['boq']), _r(u['confirmed'])
        b['fin_pct'] = _r(b['mir_value'] / b['boq_value'] * 100, 2) if b['boq_value'] > 0 else None
        b['item_coverage_pct'] = _r(b['items_with_confirmed'] / b['items'] * 100, 2) if b['items'] else None
        for k in ('boq_value', 'mir_value', 'remaining_value'): b[k] = _r(b[k], 2)
        brand_list.append(b)
    brand_list.sort(key=lambda b: -(b['boq_value'] or 0))

    valued = [x for x in rows if x['boq_value'] is not None and x['mir_value'] is not None]
    tot_boq = sum(x['boq_value'] for x in valued); tot_mir = sum(x['mir_value'] for x in valued)
    states = {}
    for rv in recs:
        states[rv['match'].get('state') or 'not_matched'] = states.get(rv['match'].get('state') or 'not_matched', 0) + 1
    summary = {'boq_total_value': _r(tot_boq, 2), 'mir_value': _r(tot_mir, 2), 'remaining_value': _r(tot_boq - tot_mir, 2),
               'fin_pct': _r(tot_mir / tot_boq * 100, 2) if tot_boq > 0 else None, 'items_valued': len(valued), 'items_unvalued': len(rows) - len(valued),
               'boq_items': len(rows), 'items_with_confirmed': sum(1 for x in rows if (x['confirmed_qty'] or 0) > 0),
               'items_pending_review': sum(1 for x in rows if x['status'] == 'pending_review'),
               'items_without_mir': sum(1 for x in rows if x['status'] == 'no_mir'),
               'items_over_boq': sum(1 for x in rows if x['status'] == 'over_boq'),
               'mir_docs': len(st['mir_docs']), 'mir_docs_ocr_done': sum(1 for d in st['mir_docs'] if d.get('ocr_status') == 'done'),
               'mir_records': len(recs), 'records_reviewed': sum(1 for r in recs if r['review_status'] == 'reviewed'),
               'records_by_match_state': states, 'unmatched_mir_records': sum(1 for r in recs if r['approved_boq_id'] is None),
               'warnings': sum(1 for e in exceptions if e['severity'] in ('warning', 'error')),
               'currency_label': st['settings'].get('currency_label') or '(not specified in workbook)'}
    sev = {'error': 0, 'warning': 1, 'info': 2}
    exceptions.sort(key=lambda e: sev.get(e['severity'], 3))
    return {'items': rows, 'brands': brand_list, 'summary': summary, 'records': recs, 'exceptions': exceptions,
            'meta': {'count_statuses': sorted(count_status), 'definition_note': DEFINITION_NOTE}}


# ----------------------------------------------------------------------------
# Excel export (formulas where the inputs are valid; "N/A" text otherwise)
# ----------------------------------------------------------------------------
def export_xlsx(st, comp, out_path, project_name=''):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter as L
    HDR = PatternFill('solid', fgColor='1F3A5F'); HF = Font(bold=True, color='FFFFFF')
    thin = Side(style='thin', color='CCCCCC'); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
    NUM, MONEY, PCT = '#,##0.##', '#,##0.00', '0.00%'
    wb = Workbook()

    def head(ws, row, titles, widths):
        for i, (t, w) in enumerate(zip(titles, widths), 1):
            c = ws.cell(row=row, column=i, value=t); c.fill, c.font, c.border = HDR, HF, BOX
            c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
            ws.column_dimensions[L(i)].width = w
        ws.row_dimensions[row].height = 32
        ws.freeze_panes = ws.cell(row=row + 1, column=1)

    # ---- Detailed (built first so other sheets can reference it)
    wd = wb.active; wd.title = 'Detailed Reconciliation'
    cols = ['Item Code (S.N)', 'BOQ Row Ref', 'Description', 'Part #', 'Brand', 'Unit', 'Original BOQ Qty', 'Recorded MIR Qty (all matched)',
            'Confirmed MIR Qty', 'Remaining Qty', 'Unit Price', 'Original BOQ Value', 'MIR Value', 'Remaining Value', 'Qty %', 'Financial %',
            'Status', 'Flags', 'Source MIR documents / pages']
    head(wd, 1, cols, [10, 16, 44, 18, 14, 8, 12, 14, 12, 12, 12, 16, 16, 16, 9, 10, 16, 34, 52])
    for n, x in enumerate(comp['items'], 2):
        src = '; '.join('%s p%s (%s, qty %g %s, status %s%s)' % (s['file'], s['page'], s['ref_no'] or 'no ref', s['qty'], s['unit'] or '', s['status'] or '?', '' if s['reviewed'] else ', UNREVIEWED') for s in x['sources'])
        vals = [x['code'], x['row_ref'], x['description'], x['part'], x['brand'], x['unit'] or 'unspecified', x['boq_qty'], x['recorded_qty'], x['confirmed_qty']]
        for i, v in enumerate(vals, 1):
            wd.cell(row=n, column=i, value=v if v not in (None, '') else 'N/A')
        qok = x['boq_qty'] is not None
        wd.cell(row=n, column=10, value='=G%d-I%d' % (n, n) if qok else 'N/A')
        pok = x['unit_price'] is not None
        wd.cell(row=n, column=11, value=x['unit_price'] if pok else 'N/A')
        wd.cell(row=n, column=12, value='=G%d*K%d' % (n, n) if (qok and pok) else 'N/A')
        wd.cell(row=n, column=13, value='=I%d*K%d' % (n, n) if pok else 'N/A')
        wd.cell(row=n, column=14, value='=L%d-M%d' % (n, n) if (qok and pok) else 'N/A')
        wd.cell(row=n, column=15, value='=I%d/G%d' % (n, n) if (qok and x['boq_qty'] > 0) else 'N/A')
        wd.cell(row=n, column=16, value='=M%d/L%d' % (n, n) if (qok and pok and x['boq_value'] and x['boq_value'] > 0) else 'N/A')
        wd.cell(row=n, column=17, value=x['status'].replace('_', ' '))
        wd.cell(row=n, column=18, value=', '.join(x['flags']))
        wd.cell(row=n, column=19, value=src)
        for c in (7, 8, 9, 10): wd.cell(row=n, column=c).number_format = NUM
        for c in (11, 12, 13, 14): wd.cell(row=n, column=c).number_format = MONEY
        for c in (15, 16): wd.cell(row=n, column=c).number_format = PCT
        fill = {'over_boq': 'F8D7DA', 'complete': 'D4EDDA', 'partial': 'FFF3CD', 'pending_review': 'FFE5B4', 'no_mir': 'EEEEEE'}.get(x['status'])
        if fill: wd.cell(row=n, column=17).fill = PatternFill('solid', fgColor=fill)
        for c in range(1, 20):
            wd.cell(row=n, column=c).border = BOX; wd.cell(row=n, column=c).alignment = Alignment(wrap_text=c in (3, 18, 19), vertical='top')
    last = max(2, len(comp['items']) + 1)
    wd.auto_filter.ref = 'A1:S%d' % last
    DR = "'Detailed Reconciliation'!"

    # ---- Executive summary
    ws = wb.create_sheet('Executive Summary', 0)
    s = comp['summary']
    ws.column_dimensions['A'].width = 38; ws.column_dimensions['B'].width = 26; ws.column_dimensions['C'].width = 70
    ws['A1'] = 'BOQ vs. MIR Reconciliation - Executive Summary'; ws['A1'].font = Font(size=15, bold=True)
    ws['A2'] = 'Project: %s' % (project_name or '-'); ws['A3'] = 'Generated (UTC): %s' % _now()
    ws['A4'] = 'BOQ file: %s (sheet "%s")' % ((st.get('boq') or {}).get('file', {}).get('filename', '-'), (st.get('boq') or {}).get('sheet', '-'))
    ws['A5'] = 'Currency: %s' % s['currency_label']
    rows = [('Original BOQ value (valued items)', "=SUM(%sL2:L%d)" % (DR, last), MONEY), ('Confirmed MIR value', "=SUM(%sM2:M%d)" % (DR, last), MONEY),
            ('Remaining BOQ value', "=SUM(%sN2:N%d)" % (DR, last), MONEY), ('Overall financial % (MIR / BOQ)', '=IFERROR(B8/B7,"N/A")', PCT),
            ('Total BOQ items', s['boq_items'], '0'), ('Items with confirmed MIR qty', s['items_with_confirmed'], '0'),
            ('Items pending review', s['items_pending_review'], '0'), ('Items with no MIR', s['items_without_mir'], '0'),
            ('Items over BOQ quantity', s['items_over_boq'], '0'), ('Items without a valid price (excluded from value)', s['items_unvalued'], '0'),
            ('MIR PDFs in reconciliation', s['mir_docs'], '0'), ('MIR PDFs with OCR run', s['mir_docs_ocr_done'], '0'),
            ('MIR records (lines)', s['mir_records'], '0'), ('MIR records reviewed', s['records_reviewed'], '0'),
            ('MIR records not matched to a BOQ row', s['unmatched_mir_records'], '0'), ('Warnings / errors', s['warnings'], '0')]
    for i, (a, b, f) in enumerate(rows, 7 - 0):
        ws.cell(row=i, column=1, value=a).font = Font(bold=True); c = ws.cell(row=i, column=2, value=b); c.number_format = f
    r0 = 7 + len(rows) + 1
    ws.cell(row=r0, column=1, value='How to read this report').font = Font(bold=True)
    ws.cell(row=r0 + 1, column=1, value=DEFINITION_NOTE).alignment = Alignment(wrap_text=True, vertical='top')
    ws.merge_cells(start_row=r0 + 1, start_column=1, end_row=r0 + 1, end_column=3); ws.row_dimensions[r0 + 1].height = 62
    r1 = r0 + 3
    ws.cell(row=r1, column=1, value='Source files (originals are preserved unchanged)').font = Font(bold=True)
    k = r1 + 1
    if st.get('boq'):
        f = st['boq']['file']; ws.cell(row=k, column=1, value='BOQ: ' + f['filename']); ws.cell(row=k, column=2, value='sha256 ' + f['sha256'][:16]); k += 1
    for d in st['mir_docs']:
        ws.cell(row=k, column=1, value='MIR: ' + d['filename']); ws.cell(row=k, column=2, value='sha256 ' + d['sha256'][:16]); ws.cell(row=k, column=3, value='%d page(s); OCR: %s' % (d['pages'], d.get('ocr_status'))); k += 1

    # ---- Brand summary (SUMIFS over the detailed sheet so it always reconciles)
    wbs = wb.create_sheet('Brand Summary', 1)
    head(wbs, 1, ['Brand', 'BOQ Items', 'Items w/ confirmed MIR', 'Item coverage %', 'Quantities by unit (BOQ / confirmed / remaining)', 'Items with unspecified unit (qty not summed)',
                  'Original BOQ Value', 'MIR Value', 'Remaining Value', 'Financial %', 'Items missing MIR', 'Items over BOQ', 'Items awaiting review', 'Unmatched MIR records'],
         [18, 9, 12, 11, 48, 16, 17, 17, 17, 10, 38, 28, 38, 12])
    for n, b in enumerate(comp['brands'], 2):
        qs = '; '.join('%s: %s / %s / %s (%s)' % (u, v['boq'], v['confirmed'], v['remaining'], ('%.1f%%' % v['pct']) if v['pct'] is not None else 'n/a') for u, v in b['qty_by_unit'].items()) or '-'
        wbs.cell(row=n, column=1, value=b['brand'])
        wbs.cell(row=n, column=2, value='=COUNTIF(%sE2:E%d,A%d)' % (DR, last, n) if b['items'] else 0)
        wbs.cell(row=n, column=3, value=b['items_with_confirmed'])
        wbs.cell(row=n, column=4, value='=IFERROR(C%d/B%d,"N/A")' % (n, n)); wbs.cell(row=n, column=4).number_format = PCT
        wbs.cell(row=n, column=5, value=qs); wbs.cell(row=n, column=6, value=b['unspecified_unit_items'])
        wbs.cell(row=n, column=7, value='=SUMIFS(%sL2:L%d,%sE2:E%d,A%d)' % (DR, last, DR, last, n))
        wbs.cell(row=n, column=8, value='=SUMIFS(%sM2:M%d,%sE2:E%d,A%d)' % (DR, last, DR, last, n))
        wbs.cell(row=n, column=9, value='=G%d-H%d' % (n, n))
        wbs.cell(row=n, column=10, value='=IF(G%d>0,H%d/G%d,"N/A")' % (n, n, n)); wbs.cell(row=n, column=10).number_format = PCT
        wbs.cell(row=n, column=11, value=', '.join(b['missing_mir']) or '-'); wbs.cell(row=n, column=12, value=', '.join(b['over_boq']) or '-')
        wbs.cell(row=n, column=13, value=', '.join(b['awaiting_review']) or '-'); wbs.cell(row=n, column=14, value=b['unmatched_mir_records'])
        for c in (7, 8, 9): wbs.cell(row=n, column=c).number_format = MONEY
        for c in range(1, 15):
            wbs.cell(row=n, column=c).border = BOX; wbs.cell(row=n, column=c).alignment = Alignment(wrap_text=True, vertical='top')
    wbs.auto_filter.ref = 'A1:N%d' % max(2, len(comp['brands']) + 1)
    tn = len(comp['brands']) + 3
    wbs.cell(row=tn, column=1, value='Total (values only)').font = Font(bold=True)
    for c, colL in ((7, 'G'), (8, 'H'), (9, 'I')):
        wbs.cell(row=tn, column=c, value='=SUM(%s2:%s%d)' % (colL, colL, tn - 2)).number_format = MONEY
    wbs.cell(row=tn + 1, column=1, value='Quantities of different units are never added together; brands are never merged.')

    # ---- MIR records
    wr = wb.create_sheet('MIR Extraction Records')
    head(wr, 1, ['Record ID', 'Source PDF', 'Page', 'Origin', 'OCR: MIR ref', 'OCR: date', 'OCR: description', 'OCR: unit', 'OCR: Q.T', 'OCR: TOTAL', 'OCR conf %',
                 'Reviewed: MIR ref', 'Reviewed: date', 'Reviewed: description', 'Reviewed: unit', 'Reviewed: qty', 'MIR status', 'Brand', 'Review status',
                 'Match state', 'Matched BOQ row', 'Match method', 'Matched/reviewed by', 'Flags'],
         [12, 26, 6, 8, 11, 11, 36, 8, 8, 8, 8, 11, 11, 36, 8, 9, 12, 12, 12, 12, 16, 12, 14, 40])
    for n, r in enumerate(comp['records'], 2):
        o, u, m = r['ocr'], r['user'], r['match']
        vals = [r['id'], r['filename'], r['page'], r['origin'], o.get('ref_no'), o.get('date'), o.get('description'), o.get('unit'), o.get('qty'), o.get('total'), o.get('conf'),
                u.get('ref_no'), u.get('date'), u.get('description'), u.get('unit'), u.get('qty'),
                (u.get('mir_status') or '') and '%s - %s' % (u['mir_status'], MIR_STATUS_LABELS[u['mir_status']]), u.get('brand'), r['review_status'],
                m.get('state') or 'not matched', r['approved_boq_id'] or m.get('boq_id') or '', m.get('method') or '', (m.get('decided_by') or r.get('edited_by') or ''), ', '.join(r['flags'])]
        for i, v in enumerate(vals, 1):
            c = wr.cell(row=n, column=i, value=v); c.border = BOX; c.alignment = Alignment(wrap_text=i in (7, 14, 24), vertical='top')
    wr.auto_filter.ref = 'A1:X%d' % max(2, len(comp['records']) + 1)

    # ---- Unmatched & review-required
    wu = wb.create_sheet('Unmatched & Review-Required')
    head(wu, 1, ['Type', 'Reference', 'Source / BOQ row', 'Description', 'Qty', 'Unit', 'Why it needs attention'], [24, 16, 30, 46, 10, 8, 70])
    n = 2
    for r in comp['records']:
        why = [f for f in r['flags'] if f not in ('manual_entry',)]
        if r['approved_boq_id'] is None or r['review_status'] != 'reviewed' or any(f.startswith('status_') or f in ('duplicate', 'unit_conflict') for f in why):
            st_ = r['match'].get('state') or 'not matched'
            wu.append(['MIR record', r['id'], '%s p%s' % (r['filename'], r['page']), r['eff']['description'], r['eff']['qty'], r['eff']['unit'],
                       'match: %s; review: %s; %s' % (st_, r['review_status'], ', '.join(why) or '-')]); n += 1
    for x in comp['items']:
        if x['status'] in ('no_mir', 'pending_review', 'over_boq', 'no_confirmed'):
            wu.append(['BOQ item - ' + x['status'].replace('_', ' '), x['row_ref'], x['brand'], x['description'], x['boq_qty'], x['unit'] or '',
                       'confirmed %s of %s; awaiting review %d; %s' % (x['confirmed_qty'], x['boq_qty'], x['awaiting_review'], ', '.join(x['flags']) or '-')])
    for row in wu.iter_rows(min_row=2):
        for c in row: c.border = BOX; c.alignment = Alignment(wrap_text=True, vertical='top')
    wu.auto_filter.ref = 'A1:G%d' % max(2, wu.max_row)

    # ---- Exceptions
    we = wb.create_sheet('Exceptions & Warnings')
    head(we, 1, ['Severity', 'Type', 'Reference', 'Message'], [10, 26, 20, 110])
    for e in comp['exceptions']:
        we.append([e['severity'], e['type'], e.get('ref', ''), e['message']])
    for row in we.iter_rows(min_row=2):
        for c in row:
            c.border = BOX; c.alignment = Alignment(wrap_text=True, vertical='top')
        sev = row[0].value
        row[0].fill = PatternFill('solid', fgColor={'error': 'F8D7DA', 'warning': 'FFF3CD'}.get(sev, 'E8F0FE'))
    we.auto_filter.ref = 'A1:D%d' % max(2, we.max_row)
    wb._sheets = [wb['Executive Summary'], wb['Brand Summary'], wb['Detailed Reconciliation'], wb['MIR Extraction Records'],
                  wb['Unmatched & Review-Required'], wb['Exceptions & Warnings']]
    wb.save(str(out_path))
    return out_path


def make_export(project_dir, st, project_name=''):
    comp = compute(st)
    out = recon_dir(project_dir) / 'exports' / ('BOQ_vs_MIR_%s.xlsx' % time.strftime('%Y%m%d-%H%M%S'))
    export_xlsx(st, comp, out, project_name)
    return out


def page_image_bytes(project_dir, st, doc_id, page, width=0):
    """Page image for the review screen. width>0 returns a smaller JPEG (fast to load inline)."""
    p = page_image_path(project_dir, st, doc_id, page)
    try: width = int(width or 0)
    except (TypeError, ValueError): width = 0
    if width > 0:
        try:
            from PIL import Image
            with Image.open(p) as im:
                width = max(300, min(width, 2000))
                im = im.convert('RGB'); im = im.resize((width, int(im.height * width / im.width)))
                buf = io.BytesIO(); im.save(buf, 'JPEG', quality=82)
                return buf.getvalue(), 'image/jpeg'
        except Exception:
            pass
    return p.read_bytes(), 'image/png'


def safe_display_name(name):
    """Display name only (stored files are always uuid-named). Keeps Arabic/Unicode letters."""
    n = Path(str(name or 'file').replace('\\', '/')).name
    n = re.sub(r'[\\/:*?"<>|\x00-\x1f]+', '_', n).strip(' .')
    return (n or 'file')[:180]


def api_payload(project_dir, st):
    """Everything the UI needs in one call."""
    comp = compute(st)
    boq = st.get('boq')
    meta = None
    if boq:
        meta = {k: boq[k] for k in ('filename', 'sheet', 'header_row', 'columns', 'warnings', 'sheets_considered', 'parsed_at', 'file')}
        meta['items'] = len(boq['items'])
    return {'ok': True, 'ocr': ocr_status(), 'boq': meta, 'docs': st['mir_docs'], 'settings': st['settings'], 'comp': comp,
            'boq_items': [{'id': it['id'], 'label': '%s | %s | qty %s | %s' % (it['id'].split('!')[-1], (it['item'] or it['reference'])[:48], it['qty'] if it['qty'] is not None else '?', it['brand'])}
                          for it in (boq or {}).get('items', [])],
            'audit': st['audit'][-60:][::-1], 'updated_at': st.get('updated_at')}
